"""
Connector and line tools for PowerPoint MCP Server.
Implements connector line/arrow drawing capabilities.
"""

from typing import Dict, List, Optional, Any
from mcp.types import ToolAnnotations
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

CONNECTOR_TYPES = {
    'straight': MSO_CONNECTOR.STRAIGHT,
    'elbow': MSO_CONNECTOR.ELBOW,
    'curved': MSO_CONNECTOR.CURVE,
}

DASH_STYLES = {
    'solid': MSO_LINE.SOLID,
    'dash': MSO_LINE.DASH,
    'dot': MSO_LINE.ROUND_DOT,
    'dash_dot': MSO_LINE.DASH_DOT,
}

# Arrowhead names accepted by the tools -> DrawingML line-end type
ARROW_TYPES = {
    'arrow': 'triangle',
    'triangle': 'triangle',
    'open_arrow': 'arrow',
    'stealth': 'stealth',
    'diamond': 'diamond',
    'oval': 'oval',
}


def _set_line_end(line, tag: str, arrow: Optional[str]) -> None:
    """Set a headEnd/tailEnd element on a line. `arrow` None or 'none' removes it."""
    ln = line._get_or_add_ln()
    for existing in ln.findall(qn(f'a:{tag}')):
        ln.remove(existing)
    if not arrow or arrow.lower() == 'none':
        return
    arrow_type = ARROW_TYPES.get(arrow.lower())
    if arrow_type is None:
        raise ValueError(f"Invalid arrow type '{arrow}'. Use: none, {', '.join(ARROW_TYPES)}")
    el = etree.SubElement(ln, qn(f'a:{tag}'))
    el.set('type', arrow_type)
    el.set('w', 'med')
    el.set('len', 'med')
    # headEnd must precede tailEnd in the schema
    if tag == 'headEnd':
        tail = ln.find(qn('a:tailEnd'))
        if tail is not None:
            ln.remove(el)
            tail.addprevious(el)


def create_connector(slide, connector_type: str, start_x: float, start_y: float,
                     end_x: float, end_y: float, line_width: float = 1.0,
                     color: Optional[List[int]] = None, arrow_end: str = "none",
                     arrow_start: str = "none", dash_style: str = "solid"):
    """Create a styled connector on `slide` (coordinates in inches) and return it.

    Raises ValueError on invalid arguments.
    """
    ctype = CONNECTOR_TYPES.get(str(connector_type).lower())
    if ctype is None:
        raise ValueError(f"Invalid connector type. Use: {list(CONNECTOR_TYPES)}")
    dash = DASH_STYLES.get(str(dash_style).lower())
    if dash is None:
        raise ValueError(f"Invalid dash_style. Use: {list(DASH_STYLES)}")
    if color is not None:
        if not (isinstance(color, list) and len(color) == 3 and
                all(isinstance(c, int) and 0 <= c <= 255 for c in color)):
            raise ValueError("color must be [r, g, b] with integers 0-255")

    connector = slide.shapes.add_connector(
        ctype, Inches(start_x), Inches(start_y), Inches(end_x), Inches(end_y)
    )
    if line_width:
        connector.line.width = Pt(line_width)
    if color:
        connector.line.color.rgb = RGBColor(*color)
    connector.line.dash_style = dash
    # Line ends go last so they land after the fill/dash children of <a:ln>.
    _set_line_end(connector.line, 'headEnd', arrow_start)
    _set_line_end(connector.line, 'tailEnd', arrow_end)
    return connector


def register_connector_tools(app, presentations, get_current_presentation_id, validate_parameters,
                          is_positive, is_non_negative, is_in_range, is_valid_rgb):
    """Register connector tools with the FastMCP app."""

    @app.tool(
        annotations=ToolAnnotations(
            title="Add Connector",
        ),
    )
    def add_connector(
        slide_index: int,
        connector_type: str,
        start_x: float,
        start_y: float,
        end_x: float,
        end_y: float,
        line_width: float = 1.0,
        color: List[int] = None,
        presentation_id: str = None,
        arrow_end: str = "none",
        arrow_start: str = "none",
        dash_style: str = "solid"
    ) -> Dict:
        """
        Add connector lines/arrows between points on a slide.

        Args:
            slide_index: Index of the slide (0-based)
            connector_type: Type of connector ("straight", "elbow", "curved")
            start_x: Starting X coordinate in inches
            start_y: Starting Y coordinate in inches
            end_x: Ending X coordinate in inches
            end_y: Ending Y coordinate in inches
            line_width: Width of the connector line in points
            color: RGB color as [r, g, b] list
            presentation_id: Optional presentation ID (uses current if not provided)
            arrow_end: Arrowhead at the end point: "none", "arrow", "open_arrow", "stealth", "diamond", "oval"
            arrow_start: Arrowhead at the start point (same options as arrow_end)
            dash_style: Line style: "solid", "dash", "dot", "dash_dot"

        Returns:
            Dictionary with operation results
        """
        try:
            pres_id = presentation_id or get_current_presentation_id()
            if pres_id not in presentations:
                return {"error": "Presentation not found"}

            pres = presentations[pres_id]

            if not (0 <= slide_index < len(pres.slides)):
                return {"error": f"Slide index {slide_index} out of range"}

            slide = pres.slides[slide_index]
            create_connector(
                slide, connector_type, start_x, start_y, end_x, end_y,
                line_width=line_width, color=color, arrow_end=arrow_end,
                arrow_start=arrow_start, dash_style=dash_style
            )

            return {
                "message": f"Added {connector_type} connector to slide {slide_index}",
                "connector_type": connector_type,
                "start_point": [start_x, start_y],
                "end_point": [end_x, end_y],
                "shape_index": len(slide.shapes) - 1
            }

        except Exception as e:
            return {"error": f"Failed to add connector: {str(e)}"}
