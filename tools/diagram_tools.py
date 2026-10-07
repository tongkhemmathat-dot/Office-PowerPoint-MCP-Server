"""
Diagram tools for PowerPoint MCP Server.
Icon library lookup (e.g. AWS architecture icons), icon+label placement and a
batch tool for building a whole diagram in one call.

Icon folders are searched recursively for png/jpg/jpeg/gif/svg files:
  - every directory listed in the PPT_ICON_PATH environment variable (os.pathsep-separated)
  - the `icons/` directory next to ppt_mcp_server.py
Run `python scripts/fetch_aws_icons.py` to populate `icons/aws/` with the AWS icon set.
"""

import hashlib
import os
import re
import tempfile
from typing import Any, Dict, List, Optional

from mcp.types import ToolAnnotations
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .connector_tools import create_connector

_IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.gif', '.svg'}
_ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}
_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _icon_dirs() -> List[str]:
    dirs = [d for d in os.environ.get("PPT_ICON_PATH", "").split(os.pathsep) if d]
    dirs.append(os.path.join(_REPO_ROOT, "icons"))
    return [d for d in dirs if os.path.isdir(d)]


def _norm(text: str) -> str:
    return re.sub(r'[^a-z0-9]', '', text.lower())


def _scan_icons() -> Dict[str, str]:
    """Map normalised icon name -> file path (first directory wins on duplicates)."""
    found: Dict[str, str] = {}
    for root_dir in _icon_dirs():
        for dirpath, _, files in os.walk(root_dir):
            for f in sorted(files):
                stem, ext = os.path.splitext(f)
                if ext.lower() in _IMAGE_EXTS:
                    found.setdefault(_norm(stem), os.path.join(dirpath, f))
    return found


def _resolve_icon(name: str) -> str:
    """Return a PNG/JPG path for an icon name, a file path, or raise ValueError."""
    if os.path.isfile(name):
        path = name
    else:
        icons = _scan_icons()
        key = _norm(os.path.splitext(os.path.basename(name))[0])
        if not icons:
            raise ValueError(
                "No icon library found. Set PPT_ICON_PATH or run scripts/fetch_aws_icons.py "
                "(searched: " + (", ".join(_icon_dirs()) or "no existing icon directories") + ")")
        if key in icons:
            path = icons[key]
        else:
            matches = sorted(k for k in icons if key and key in k)
            if len(matches) == 1:
                path = icons[matches[0]]
            elif matches:
                raise ValueError(f"Icon '{name}' is ambiguous: {', '.join(matches[:10])}")
            else:
                raise ValueError(f"Icon '{name}' not found. Use list_icons to search.")
    if path.lower().endswith('.svg'):
        path = _svg_to_png(path)
    return path


def _svg_to_png(svg_path: str) -> str:
    try:
        import cairosvg
    except ImportError:
        raise ValueError("SVG icons need the 'cairosvg' package (pip install cairosvg), "
                         "or use a PNG version of the icon")
    digest = hashlib.md5(os.path.abspath(svg_path).encode()).hexdigest()[:10]
    out = os.path.join(tempfile.gettempdir(), f"ppt_icon_{digest}.png")
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(svg_path):
        cairosvg.svg2png(url=svg_path, write_to=out, output_width=512)
    return out


def _check_rgb(color, field: str):
    if color is None:
        return None
    if not (isinstance(color, list) and len(color) == 3 and
            all(isinstance(c, int) and 0 <= c <= 255 for c in color)):
        raise ValueError(f"{field} must be [r, g, b] with integers 0-255")
    return RGBColor(*color)


def _add_text(slide, text: str, left: float, top: float, width: float, height: float,
              font_size: int = 12, bold: bool = False, color: Optional[List[int]] = None,
              alignment: str = "left", font_name: Optional[str] = None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    rgb = _check_rgb(color, "color")
    for i, line in enumerate(str(text).split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = _ALIGN.get(alignment, PP_ALIGN.LEFT)
        run = para.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if rgb is not None:
            run.font.color.rgb = rgb
        if font_name:
            run.font.name = font_name
    return box


def _add_icon(slide, name: str, left: float, top: float, size: float = 0.6,
              label: Optional[str] = None, label_font_size: int = 9,
              label_color: Optional[List[int]] = None, label_width: Optional[float] = None,
              label_height: float = 0.4) -> List[int]:
    """Place an icon (and optional centred label underneath). Returns new shape indexes."""
    path = _resolve_icon(name)
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(size), Inches(size))
    indexes = [len(slide.shapes) - 1]
    if label:
        lw = label_width or max(size * 2, 1.2)
        _add_text(slide, label, left + size / 2 - lw / 2, top + size + 0.02, lw, label_height,
                  font_size=label_font_size, color=label_color, alignment="center")
        indexes.append(len(slide.shapes) - 1)
    return indexes


def _add_shape(slide, add_shape_direct, shape_type: str, left: float, top: float,
               width: float, height: float, fill_color=None, line_color=None,
               line_width: Optional[float] = None, text: Optional[str] = None,
               font_size: Optional[int] = None, font_color=None):
    fill, line, font = (_check_rgb(fill_color, "fill_color"), _check_rgb(line_color, "line_color"),
                        _check_rgb(font_color, "font_color"))
    shape = add_shape_direct(slide, shape_type, left, top, width, height)
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is not None:
        shape.line.color.rgb = line
    if line_width is not None:
        shape.line.width = Pt(line_width)
    if text:
        shape.text_frame.text = text
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if font_size:
                    run.font.size = Pt(font_size)
                if font is not None:
                    run.font.color.rgb = font
    return shape


def register_diagram_tools(app, presentations, get_current_presentation_id, add_shape_direct):
    """Register icon and diagram tools with the FastMCP app."""

    def _get_slide(slide_index: int, presentation_id: Optional[str]):
        pres_id = presentation_id or get_current_presentation_id()
        if pres_id not in presentations:
            raise ValueError("Presentation not found")
        pres = presentations[pres_id]
        if not (0 <= slide_index < len(pres.slides)):
            raise ValueError(f"Slide index {slide_index} out of range")
        return pres.slides[slide_index]

    @app.tool(
        annotations=ToolAnnotations(
            title="List Icons",
            readOnlyHint=True,
        ),
    )
    def list_icons(query: Optional[str] = None, limit: int = 50) -> Dict:
        """
        Search the icon library (folders from PPT_ICON_PATH and the repo's icons/ directory).

        Args:
            query: Case-insensitive text the icon name must contain (e.g. "ec2", "load")
            limit: Maximum number of icons to return
        """
        icons = _scan_icons()
        key = _norm(query) if query else ""
        names = sorted(k for k in icons if key in k)
        return {
            "search_directories": _icon_dirs(),
            "total_icons": len(icons),
            "matches": len(names),
            "icons": [{"name": n, "path": icons[n]} for n in names[:max(limit, 0)]],
        }

    @app.tool(
        annotations=ToolAnnotations(
            title="Add Icon",
        ),
    )
    def add_icon(
        slide_index: int,
        name: str,
        left: float,
        top: float,
        size: float = 0.6,
        label: Optional[str] = None,
        label_font_size: int = 9,
        label_color: Optional[List[int]] = None,
        label_width: Optional[float] = None,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """
        Add an icon from the icon library, with an optional centred label underneath.

        Args:
            slide_index: Index of the slide (0-based)
            name: Icon name (see list_icons; case/punctuation-insensitive, unique partial
                names work) or a path to an image file
            left: Left position in inches
            top: Top position in inches
            size: Icon width and height in inches
            label: Text under the icon; use \\n for multiple lines
            label_font_size: Label font size in points
            label_color: Label RGB color as [r, g, b]
            label_width: Label width in inches (default: max(2*size, 1.2))
            presentation_id: Optional presentation ID (uses current if not provided)
        """
        try:
            slide = _get_slide(slide_index, presentation_id)
            indexes = _add_icon(slide, name, left, top, size, label, label_font_size,
                                label_color, label_width)
            return {"message": f"Added icon '{name}' to slide {slide_index}",
                    "icon_shape_index": indexes[0],
                    "label_shape_index": indexes[1] if len(indexes) > 1 else None}
        except Exception as e:
            return {"error": f"Failed to add icon: {str(e)}"}

    @app.tool(
        annotations=ToolAnnotations(
            title="Add Diagram Elements",
        ),
    )
    def add_diagram_elements(
        slide_index: int,
        elements: List[Dict[str, Any]],
        presentation_id: Optional[str] = None
    ) -> Dict:
        """
        Add many elements to a slide in one call. Elements are added in order, so earlier
        elements sit behind later ones (put containers like VPC boxes first).

        Each element is an object with a "type" and that type's fields (inches / points):
          shape:     shape_type, left, top, width, height; optional fill_color, line_color,
                     line_width, text, font_size, font_color
          icon:      name, left, top; optional size, label, label_font_size, label_color, label_width
          text:      text, left, top, width, height; optional font_size, bold, color,
                     alignment (left|center|right), font_name
          connector: start_x, start_y, end_x, end_y; optional connector_type (straight|elbow|curved),
                     line_width, color, arrow_end, arrow_start, dash_style

        Example:
          [{"type": "shape", "shape_type": "rectangle", "left": 1, "top": 1, "width": 4, "height": 3,
            "fill_color": [240, 248, 225], "line_color": [122, 161, 22]},
           {"type": "icon", "name": "EC2", "left": 2, "top": 1.5, "label": "Web server"},
           {"type": "connector", "start_x": 2.6, "start_y": 1.8, "end_x": 5, "end_y": 1.8, "arrow_end": "arrow"}]

        Invalid elements are skipped and reported in `errors` (with their position in the list);
        valid ones are still added.

        Args:
            slide_index: Index of the slide (0-based)
            elements: List of element objects described above
            presentation_id: Optional presentation ID (uses current if not provided)
        """
        try:
            slide = _get_slide(slide_index, presentation_id)
        except Exception as e:
            return {"error": f"Failed to add elements: {str(e)}"}

        added: List[Dict[str, Any]] = []
        errors: List[Dict[str, Any]] = []
        for pos, el in enumerate(elements):
            try:
                if not isinstance(el, dict):
                    raise ValueError("element must be an object")
                kind = str(el.get("type", "")).lower()
                f = {k: v for k, v in el.items() if k != "type"}
                if kind == "shape":
                    _add_shape(slide, add_shape_direct, **f)
                    added.append({"element": pos, "type": kind, "shape_index": len(slide.shapes) - 1})
                elif kind == "icon":
                    idx = _add_icon(slide, **f)
                    added.append({"element": pos, "type": kind, "shape_index": idx[0],
                                  "label_shape_index": idx[1] if len(idx) > 1 else None})
                elif kind == "text":
                    _add_text(slide, **f)
                    added.append({"element": pos, "type": kind, "shape_index": len(slide.shapes) - 1})
                elif kind == "connector":
                    f.setdefault("connector_type", "straight")
                    create_connector(slide, **f)
                    added.append({"element": pos, "type": kind, "shape_index": len(slide.shapes) - 1})
                else:
                    raise ValueError(f"unknown type '{el.get('type')}'. Use shape, icon, text or connector")
            except Exception as e:  # includes TypeError for unexpected/missing field names
                errors.append({"element": pos, "error": str(e)})

        result: Dict[str, Any] = {
            "message": f"Added {len(added)} of {len(elements)} elements to slide {slide_index}",
            "added": added,
        }
        if errors:
            result["errors"] = errors
        return result
