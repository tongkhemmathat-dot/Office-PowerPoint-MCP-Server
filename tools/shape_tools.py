"""
Shape editing tools for PowerPoint MCP Server.
List, move, resize, delete and re-order shapes that already exist on a slide.
"""

from typing import Dict, Optional
from mcp.types import ToolAnnotations
from pptx.util import Inches, Emu

# Children of <p:spTree> that come before the first drawable shape.
_SPTREE_HEADER_TAGS = ('nvGrpSpPr', 'grpSpPr')


def _to_inches(emu) -> Optional[float]:
    return None if emu is None else round(Emu(emu).inches, 3)


def describe_shape(index: int, shape) -> Dict:
    """Summarise a shape for list output."""
    info = {
        "shape_index": index,
        "shape_id": shape.shape_id,
        "name": shape.name,
        "type": str(shape.shape_type).split('.')[-1].split(' ')[0] if shape.shape_type else "UNKNOWN",
        "left": _to_inches(shape.left),
        "top": _to_inches(shape.top),
        "width": _to_inches(shape.width),
        "height": _to_inches(shape.height),
    }
    if shape.has_text_frame and shape.text_frame.text:
        text = shape.text_frame.text
        info["text"] = text if len(text) <= 60 else text[:57] + "..."
    return info


def register_shape_tools(app, presentations, get_current_presentation_id):
    """Register shape editing tools with the FastMCP app."""

    @app.tool(
        annotations=ToolAnnotations(
            title="Manage Shapes",
            destructiveHint=True,
        ),
    )
    def manage_shapes(
        slide_index: int,
        operation: str,
        shape_index: Optional[int] = None,
        left: Optional[float] = None,
        top: Optional[float] = None,
        width: Optional[float] = None,
        height: Optional[float] = None,
        dx: Optional[float] = None,
        dy: Optional[float] = None,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """
        List and edit existing shapes on a slide.

        Operations:
            list: List every shape (index, id, name, type, position, size, text).
            move: Move a shape. Give absolute `left`/`top` and/or relative `dx`/`dy` (inches).
            resize: Set `width` and/or `height` (inches).
            delete: Delete the shape. Later shapes shift down by one index, so re-list afterwards.
            bring_to_front: Move the shape to the top of the z-order.
            send_to_back: Move the shape to the bottom of the z-order.

        Args:
            slide_index: Index of the slide (0-based)
            operation: One of list, move, resize, delete, bring_to_front, send_to_back
            shape_index: Index of the target shape (from `list` or from the tool that created it)
            left, top, width, height: Absolute values in inches
            dx, dy: Relative offsets in inches (move only)
            presentation_id: Optional presentation ID (uses current if not provided)
        """
        try:
            pres_id = presentation_id or get_current_presentation_id()
            if pres_id not in presentations:
                return {"error": "Presentation not found"}
            pres = presentations[pres_id]
            if not (0 <= slide_index < len(pres.slides)):
                return {"error": f"Slide index {slide_index} out of range"}
            slide = pres.slides[slide_index]
            shapes = list(slide.shapes)
            op = operation.lower()

            if op == "list":
                return {
                    "slide_index": slide_index,
                    "shape_count": len(shapes),
                    "shapes": [describe_shape(i, s) for i, s in enumerate(shapes)],
                }

            if op not in ("move", "resize", "delete", "bring_to_front", "send_to_back"):
                return {"error": "Invalid operation. Use: list, move, resize, delete, bring_to_front, send_to_back"}
            if shape_index is None or not (0 <= shape_index < len(shapes)):
                return {"error": f"shape_index {shape_index} out of range (slide has {len(shapes)} shapes)"}
            shape = shapes[shape_index]

            if op == "move":
                if all(v is None for v in (left, top, dx, dy)):
                    return {"error": "move needs left/top and/or dx/dy"}
                if left is not None:
                    shape.left = Inches(left)
                if top is not None:
                    shape.top = Inches(top)
                if dx is not None:
                    shape.left = shape.left + Inches(dx)
                if dy is not None:
                    shape.top = shape.top + Inches(dy)
            elif op == "resize":
                if width is None and height is None:
                    return {"error": "resize needs width and/or height"}
                if width is not None:
                    shape.width = Inches(width)
                if height is not None:
                    shape.height = Inches(height)
            elif op == "delete":
                el = shape._element
                el.getparent().remove(el)
                return {"message": f"Deleted shape {shape_index} ('{shape.name}')",
                        "shape_count": len(slide.shapes)}
            else:
                el = shape._element
                tree = el.getparent()
                tree.remove(el)
                if op == "bring_to_front":
                    tree.append(el)
                else:
                    pos = sum(1 for c in tree if c.tag.split('}')[-1] in _SPTREE_HEADER_TAGS)
                    tree.insert(pos, el)

            new_index = list(slide.shapes).index(shape) if op in ("bring_to_front", "send_to_back") else shape_index
            return {"message": f"{op} applied to shape {shape_index} ('{shape.name}')",
                    "shape": describe_shape(new_index, shape)}

        except Exception as e:
            return {"error": f"Failed to {operation} shape: {str(e)}"}
